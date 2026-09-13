"""Pra-proses: dari arsip voltamogram mentah ke tiga berkas npz yang dibaca
calpir.py.

  scans_LC.npz, scans_EM.npz  sepuluh sapuan tiap berkas diresample ke 181 titik:
                              V dan I [n_berkas, 10, 181], y dalam ppm, ok
  sumbu_sama.npz              sumbu potensial kedua alat disamakan, yaitu cabang
                              katodik dan anodik tiap sapuan diinterpolasi ke kisi
                              -1,40 sampai -0,50 V langkah 0,01 V, lalu median
                              sapuan 3 sampai 10, 178 peubah bersama

Masukan adalah 01_raw_voltammograms.zip dari Zenodo, doi 10.5281/zenodo.22236054.
Keluaran ditulis ke data/, atau ke direktori CALPIR_DATA.

Sapuan LecSens berjalan menyambung, sehingga titik ke-i kedua alat tidak berada
pada potensial yang sama. Karena itu tiap sapuan dipecah menurut arah geraknya
dan diinterpolasi terhadap potensial, bukan dibandingkan per indeks.

Pemakaian:
    python3 pra_proses.py /jalur/ke/01_raw_voltammograms.zip
"""
import io
import os
import sys
import zipfile

import numpy as np
import openpyxl

from calpir import DATA

FOLDER = {"LC": "Alat Sensor Cadmium LC", "EM": "Alat Sensor CD Groundturth Estat"}
LEVEL_PPM = [2, 4, 6, 8, 10, 20, 40, 60, 80, 100, 200, 400, 600, 800, 1000]
N_TITIK, N_SAPUAN = 181, 10
KISI = np.round(np.arange(-1.40, -0.50 + 0.01 / 2, 0.01), 4)
BUANG_SAPUAN = 2          # dua sapuan pertama transien kondisioning
MIN_TITIK = 15            # sebuah arah dipakai bila punya cukup titik
AMBANG_KOSONG = 0.02      # peubah dibuang bila kosong pada lebih dari 2 persen berkas


# ------------------------------------------------------ baca arsip mentah
def _satu_sapuan_mentah(baris, s):
    """Kolom 2s berisi potensial dan kolom 2s+1 berisi arus sapuan ke-s."""
    pasangan = [(r[2 * s], r[2 * s + 1]) for r in baris
                if r[2 * s] is not None and r[2 * s + 1] is not None]
    if not pasangan:
        return None, None
    return (np.array([p[0] for p in pasangan], dtype=float),
            np.array([p[1] for p in pasangan], dtype=float))


def _resample(v, i):
    """Panjang sapuan berbeda 164 sampai 183 titik, diresample ke 181 terhadap
    indeks posisi, sebab potensial tidak monoton sepanjang satu siklus."""
    asal, tuju = np.linspace(0.0, 1.0, num=len(i)), np.linspace(0.0, 1.0, num=N_TITIK)
    return np.interp(tuju, asal, v), np.interp(tuju, asal, i)


def baca_arsip(zf, alat):
    V, I, y, ok = [], [], [], []
    for level in LEVEL_PPM:
        awalan = f"{FOLDER[alat]}/{level}/"
        berkas = sorted(n for n in zf.namelist() if n.startswith(awalan) and n.endswith(".xlsx"))
        for f in berkas:
            wb = openpyxl.load_workbook(io.BytesIO(zf.read(f)), data_only=True)
            baris = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[2:]
            vs, is_ = np.zeros((N_SAPUAN, N_TITIK)), np.zeros((N_SAPUAN, N_TITIK))
            bendera = np.zeros(N_SAPUAN, dtype=bool)
            for s in range(N_SAPUAN):
                v, i = _satu_sapuan_mentah(baris, s)
                if v is None or len(v) < 20:
                    continue
                vs[s], is_[s] = _resample(v, i)
                bendera[s] = True
            V.append(vs); I.append(is_); y.append(level); ok.append(bendera)
        print(f"  {alat} {level:>4} ppm: {len(berkas)} berkas", flush=True)
    return np.stack(V), np.stack(I), np.array(y), np.stack(ok)


# ---------------------------------------------- samakan sumbu potensial
def _satu_arah(v, i, naik):
    """Titik dikumpulkan dari seluruh ruas yang arahnya sama, sebab sapuan
    LecSens dapat terpotong menjadi dua ruas oleh pembalikan arah."""
    arah = np.sign(np.diff(v))
    pilih = arah > 0 if naik else arah < 0
    if pilih.sum() < MIN_TITIK:
        return np.full(KISI.size, np.nan)
    idx = np.zeros(v.size, dtype=bool)
    idx[:-1] |= pilih
    idx[1:] |= pilih
    vv, ii = v[idx], i[idx]
    urut = np.argsort(vv)
    vv, ii = vv[urut], ii[urut]
    vv, awal = np.unique(np.round(vv, 6), return_inverse=True)
    ii = np.bincount(awal, weights=ii) / np.bincount(awal)
    return np.interp(KISI, vv, ii, left=np.nan, right=np.nan)


def satu_berkas(V, I, ok):
    """Median sapuan 3 sampai 10, masing-masing katodik lalu anodik."""
    baris = [np.concatenate([_satu_arah(V[s], I[s], False), _satu_arah(V[s], I[s], True)])
             for s in range(BUANG_SAPUAN, V.shape[0]) if ok[s]]
    if not baris:
        return np.full(2 * KISI.size, np.nan)
    with np.errstate(all="ignore"):
        return np.nanmedian(np.vstack(baris), axis=0)


def tambal(*matriks):
    """Peubah dibuang bersama pada kedua alat agar sumbunya identik; sisa
    kekosongan kecil ditambal median kolomnya."""
    layak = np.ones(matriks[0].shape[1], dtype=bool)
    for X in matriks:
        layak &= np.isnan(X).mean(axis=0) <= AMBANG_KOSONG
    keluar = []
    for X in matriks:
        Z = X[:, layak].copy()
        med = np.nanmedian(Z, axis=0)
        b, k = np.where(np.isnan(Z))
        Z[b, k] = med[k]
        keluar.append(Z)
    return keluar, layak


def main():
    arsip = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("CALPIR_ZIP", "")
    assert os.path.isfile(arsip), "tunjuk 01_raw_voltammograms.zip sebagai argumen pertama"
    os.makedirs(DATA, exist_ok=True)
    X, y = {}, {}
    with zipfile.ZipFile(arsip) as zf:
        for alat in ("LC", "EM"):
            V, I, y[alat], ok = baca_arsip(zf, alat)
            assert V.shape[0] > 0, f"tidak ada berkas {alat}, periksa isi arsip"
            np.savez_compressed(os.path.join(DATA, f"scans_{alat}.npz"), V=V, I=I, y=y[alat], ok=ok)
            X[alat] = np.vstack([satu_berkas(V[n], I[n], ok[n]) for n in range(V.shape[0])])
            print(f"{alat}: {V.shape[0]} berkas, sapuan terbaca {ok.sum()}/{ok.size}", flush=True)
    (X_LC, X_EM), layak = tambal(X["LC"], X["EM"])
    sumbu = np.concatenate([KISI, KISI])[layak]
    cabang = np.concatenate([np.zeros(KISI.size), np.ones(KISI.size)])[layak]
    np.savez_compressed(os.path.join(DATA, "sumbu_sama.npz"), X_LC=X_LC, y_LC=y["LC"],
                        X_EM=X_EM, y_EM=y["EM"], sumbu=sumbu, cabang=cabang)
    print(f"peubah bersama {layak.sum()} dari {layak.size}, ditulis ke {DATA}")


if __name__ == "__main__":
    main()
