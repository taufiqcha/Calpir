"""Batas deteksi, batas kuantifikasi, dan perilaku pada elektrolit blanko.

Reviewer 1 butir 2 meminta batas deteksi dan batas kuantifikasi dilaporkan.
Reviewer 2 butir 3 menyatakan menyembunyikan satu tingkat kadar belum merupakan
uji luar distribusi yang kuat, sebab analit, elektroda, elektrolit, dan sel
tetap sama.

Blanko menjawab keduanya sekaligus. Lima puluh replikat elektrolit tanpa
kadmium per alat memberi simpangan baku blanko yang diperlukan kaidah tiga
sigma, dan sekaligus merupakan sampel yang kimianya memang berbeda, bukan
sekadar satu tingkat yang dicoret dari rentang.

Skrip ini mengerjakan empat hal.

  1. Membaca replikat blanko dari arsip mentah, lalu menempatkannya pada sumbu
     potensial yang sama persis dengan data terukur.
  2. Menghitung batas deteksi dan batas kuantifikasi dari arus puncak katodik,
     dengan kaidah tiga sigma dan sepuluh sigma terhadap kemiringan tanggapan.
     Kecocokan garisnya ikut dilaporkan apa adanya.
  3. Menjalankan aturan penolakan dua tahap dengan blanko sebagai himpunan uji.
  4. Menjalankan selang CALPIR pada blanko, lalu mencatat kadar yang dilaporkan.

Keluaran ditulis ke results/hasil_lod_blanko.json.

Cara pakai: python3 uji_lod.py [berkas_arsip.zip]
"""
import io
import json
import os
import sys
import time
import zipfile

import numpy as np
import openpyxl
from sklearn.ensemble import RandomForestRegressor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pra_proses
from calpir import (ALAT, ALPHA, BETA, BUANG_GERBANG, DATA, HASIL, ambang,
                    belah, muat, penjaga_dua_tahap, skala_sebar)
from pra_proses import FOLDER, KISI, N_SAPUAN, N_TITIK, satu_berkas

N_BENIH = 20
RENTANG_RENDAH = 20.0     # batas atas ruas yang diuji kelurusannya


def baca_blanko(zf, alat):
    """Replikat elektrolit blanko satu alat, dibaca seperti data terukur."""
    awalan = f"{FOLDER[alat]}/KCL/"
    berkas = sorted(n for n in zf.namelist()
                    if n.startswith(awalan) and n.endswith(".xlsx"))
    V, I, ok = [], [], []
    for f in berkas:
        wb = openpyxl.load_workbook(io.BytesIO(zf.read(f)), data_only=True)
        baris = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[2:]
        vs, is_ = np.zeros((N_SAPUAN, N_TITIK)), np.zeros((N_SAPUAN, N_TITIK))
        bendera = np.zeros(N_SAPUAN, dtype=bool)
        for s in range(N_SAPUAN):
            v, i = pra_proses._satu_sapuan_mentah(baris, s)
            if v is None or len(v) < 20:
                continue
            vs[s], is_[s] = pra_proses._resample(v, i)
            bendera[s] = True
        V.append(vs); I.append(is_); ok.append(bendera)
    print(f"  {alat} blanko: {len(berkas)} berkas", flush=True)
    return np.stack(V), np.stack(I), np.stack(ok)


def topeng_kolom():
    """Susun ulang kolom yang dipertahankan, dari sumbu dan cabang tersimpan."""
    z = np.load(os.path.join(DATA, "sumbu_sama.npz"))
    penuh = list(zip(np.concatenate([KISI, KISI]).round(4),
                     np.concatenate([np.zeros(KISI.size), np.ones(KISI.size)])))
    simpan = set(zip(z["sumbu"].round(4), z["cabang"]))
    return np.array([p in simpan for p in penuh], bool)


def puncak(V, I, ok):
    """Arus puncak katodik tiap replikat, median atas sapuan yang terbaca."""
    keluar = []
    for n in range(V.shape[0]):
        nilai = [I[n, s].min() for s in range(pra_proses.BUANG_SAPUAN, V.shape[1])
                 if ok[n, s]]
        keluar.append(float(np.median(nilai)) if nilai else np.nan)
    return np.array(keluar)


def titik_terbaik(X, y):
    """Titik potensial yang tanggapannya paling erat mengikuti kadar.

    Arus puncak katodik memuat arus kapasitif yang tidak membawa kadar, sehingga
    kalibrasi univariat atasnya lemah. Memilih satu titik potensial yang paling
    erat berperingkat dengan kadar meniru praktik memilih panjang gelombang
    analitik, dan memberi batas deteksi yang lebih adil bagi platformnya.
    """
    from scipy.stats import spearmanr
    rho = np.array([abs(spearmanr(X[:, j], y).statistic) for j in range(X.shape[1])])
    rho[~np.isfinite(rho)] = 0.0
    return int(np.argmax(rho)), float(rho.max())


def lod_loq(kadar, respons, sigma, batas_atas=None):
    """Kaidah tiga sigma dan sepuluh sigma terhadap kemiringan garis kalibrasi."""
    m = np.ones(kadar.size, bool) if batas_atas is None else kadar <= batas_atas
    if m.sum() < 3:
        return None
    k, r = kadar[m], respons[m]
    kem, pot = np.polyfit(k, r, 1)
    sisa = r - (kem * k + pot)
    r2 = 1 - sisa.var() / r.var() if r.var() > 0 else 0.0
    if kem == 0:
        return None
    return {"kemiringan": float(kem), "potongan": float(pot), "R2": float(r2),
            "n_titik": int(m.sum()),
            "LOD_ppm": float(3 * sigma / abs(kem)),
            "LOQ_ppm": float(10 * sigma / abs(kem))}


def uji_blanko(X, y, t, level, Xb):
    """Penolakan dan kadar yang dilaporkan pada blanko, atas 20 benih."""
    tolak, kadar_lapor, tutup_bawah = [], [], []
    for b in range(N_BENIH):
        rng = np.random.default_rng(1000 + b)
        i_fit, i_kal, i_uji = belah(y, level, rng, 20, 20, 10)
        tolak.append(penjaga_dua_tahap(X[i_fit], X[i_kal], Xb,
                                       BETA, BUANG_GERBANG)[0])
        rf = RandomForestRegressor(n_estimators=300, random_state=b,
                                   n_jobs=-1).fit(X[i_fit], t[i_fit])
        p_kal, p_b = rf.predict(X[i_kal]), rf.predict(Xb)
        s_kal = skala_sebar(rf, X, i_kal, lantai=1e-6)
        s_b = np.std([e.predict(Xb) for e in rf.estimators_], axis=0) + 1e-6
        q = ambang(np.abs(t[i_kal] - p_kal) / s_kal, ALPHA)
        kadar_lapor.append(float(np.median(10 ** p_b)))
        tutup_bawah.append(float(np.median(10 ** (p_b - q * s_b))))
        print(f"  benih {b} selesai", flush=True)
    return {"tolak_blanko_rata": float(np.mean(tolak)),
            "tolak_blanko_sd": float(np.std(tolak, ddof=1)),
            "tolak_blanko_min": float(np.min(tolak)),
            "kadar_dilaporkan_median_ppm": float(np.median(kadar_lapor)),
            "batas_bawah_median_ppm": float(np.median(tutup_bawah)),
            "n_benih": N_BENIH}


def main(arsip):
    data = muat()
    layak = topeng_kolom()
    zs = np.load(os.path.join(DATA, "sumbu_sama.npz"))
    z_sumbu, z_cabang = zs["sumbu"], zs["cabang"]
    keluar = {}
    with zipfile.ZipFile(arsip) as zf:
        for alat in ALAT:
            print(alat, flush=True)
            Vb, Ib, okb = baca_blanko(zf, alat)
            Xb_penuh = np.vstack([satu_berkas(Vb[n], Ib[n], okb[n])
                                  for n in range(Vb.shape[0])])
            X, y, t, level = data[alat]
            Xb = Xb_penuh[:, layak]
            med = np.nanmedian(X, axis=0)
            bar, kol = np.where(np.isnan(Xb))
            Xb[bar, kol] = med[kol]
            assert Xb.shape[1] == X.shape[1], "sumbu blanko tidak sepadan"

            z = np.load(os.path.join(DATA, f"scans_{alat}.npz"))
            p_level = puncak(z["V"], z["I"], z["ok"])
            p_blanko = puncak(Vb, Ib, okb)
            sigma = float(np.std(p_blanko, ddof=1))
            kadar_rata = np.array([float(np.nanmean(p_level[z["y"] == lv]))
                                   for lv in level])

            keluar[alat] = {
                "n_blanko": int(Xb.shape[0]),
                "puncak_blanko": {"rata": float(np.mean(p_blanko)), "sd": sigma,
                                  "min": float(p_blanko.min()),
                                  "maks": float(p_blanko.max())},
                "puncak_per_level": {str(int(lv)): float(v)
                                     for lv, v in zip(level, kadar_rata)},
                "lod_rentang_penuh": lod_loq(level.astype(float), kadar_rata, sigma),
                "lod_rentang_rendah": lod_loq(level.astype(float), kadar_rata,
                                              sigma, RENTANG_RENDAH),
            }

            j, rho = titik_terbaik(X, y)
            r_level = np.array([float(X[y == lv, j].mean()) for lv in level])
            sigma_j = float(np.std(Xb[:, j], ddof=1))
            keluar[alat]["titik_terpilih"] = {
                "indeks": j, "potensial_V": float(z_sumbu[j]),
                "cabang": "katodik" if z_cabang[j] == 0 else "anodik",
                "spearman_mutlak": rho, "sd_blanko": sigma_j}
            keluar[alat]["lod_titik_penuh"] = lod_loq(level.astype(float), r_level, sigma_j)
            keluar[alat]["lod_titik_rendah"] = lod_loq(level.astype(float), r_level,
                                                       sigma_j, RENTANG_RENDAH)
            keluar[alat]["blanko_sebagai_luar_distribusi"] = uji_blanko(X, y, t, level, Xb)
    keluar["catatan"] = {"alpha": ALPHA, "beta": BETA,
                         "buang_gerbang": BUANG_GERBANG,
                         "rentang_rendah_ppm": RENTANG_RENDAH,
                         "kaidah": "LOD = 3 sigma blanko / kemiringan, LOQ = 10 sigma / kemiringan",
                         "arsip": os.path.basename(arsip)}
    p = os.path.join(HASIL, "hasil_lod_blanko.json")
    json.dump(keluar, open(p, "w"), indent=1)
    print("ditulis", p, flush=True)

    for alat in ALAT:
        h = keluar[alat]
        print(f"\n{alat}  {h['n_blanko']} replikat blanko, "
              f"sigma {h['puncak_blanko']['sd']:.3f}")
        t = h["titik_terpilih"]
        print(f"  titik terpilih {t['potensial_V']:.2f} V {t['cabang']}, "
              f"|rho|={t['spearman_mutlak']:.3f}, sd blanko {t['sd_blanko']:.3f}")
        for nm in ("lod_rentang_penuh", "lod_rentang_rendah",
                   "lod_titik_penuh", "lod_titik_rendah"):
            v = h[nm]
            if v is None:
                print(f"  {nm}: tidak terhitung")
                continue
            print(f"  {nm:20s} n={v['n_titik']:2d}  R2={v['R2']:.3f}  "
                  f"LOD={v['LOD_ppm']:.2f} ppm  LOQ={v['LOQ_ppm']:.2f} ppm")
        b = h["blanko_sebagai_luar_distribusi"]
        print(f"  blanko ditolak {b['tolak_blanko_rata']:.3f} "
              f"(sd {b['tolak_blanko_sd']:.3f}, min {b['tolak_blanko_min']:.3f}), "
              f"kadar dilaporkan {b['kadar_dilaporkan_median_ppm']:.2f} ppm")


if __name__ == "__main__":
    t0 = time.time()
    arsip = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("CALPIR_ZIP", "")
    assert os.path.isfile(arsip), "tunjuk arsip voltamogram mentah sebagai argumen"
    main(arsip)
    print(f"selesai dalam {(time.time() - t0) / 60:.1f} menit", flush=True)
